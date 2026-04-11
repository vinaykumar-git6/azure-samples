"""
Run complete replication test: write to primary and read from secondary in parallel
"""
import subprocess
import threading
import time
import os
import sys
import json
from dotenv import load_dotenv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

load_dotenv()

def run_writer():
    """Run writer in a separate thread"""
    print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] Starting writer thread...")
    result = subprocess.run([sys.executable, 'redis_writer.py'], capture_output=False)
    if result.returncode != 0:
        print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] ERROR: Writer failed!")
    else:
        print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] ✓ Writer completed")
    return result.returncode

def run_reader():
    """Run reader in a separate thread"""
    # Give writer a small head start to begin writing
    retry_delay = int(os.getenv('RETRY_DELAY_SECONDS', '5'))
    print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] Reader waiting {retry_delay} seconds before starting...")
    time.sleep(retry_delay)
    
    print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] Starting reader thread...")
    result = subprocess.run([sys.executable, 'redis_reader.py'], capture_output=False)
    if result.returncode != 0:
        print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] ERROR: Reader failed!")
    else:
        print(f"[{datetime.utcnow().strftime('%H:%M:%S.%f')[:-3]}] ✓ Reader completed")
    return result.returncode

def run_replication_test():
    """Run parallel write and read test"""
    
    print("=" * 80)
    print("Redis Geo-Replication Test (Parallel Mode)")
    print("=" * 80)
    print("This test runs writer and reader in parallel to measure real-time replication.")
    print("Writer starts immediately, reader starts after a short delay.")
    print("=" * 80)
    
    # Track return codes
    writer_result = None
    reader_result = None
    
    def writer_thread():
        nonlocal writer_result
        writer_result = run_writer()
    
    def reader_thread():
        nonlocal reader_result
        reader_result = run_reader()
    
    # Start both threads
    start_time = datetime.utcnow()
    print(f"\n[{start_time.strftime('%H:%M:%S.%f')[:-3]}] Starting parallel test...")
    print("=" * 80)
    
    writer_t = threading.Thread(target=writer_thread, name="WriterThread")
    reader_t = threading.Thread(target=reader_thread, name="ReaderThread")
    
    writer_t.start()
    reader_t.start()
    
    # Wait for both to complete
    writer_t.join()
    reader_t.join()
    
    end_time = datetime.utcnow()
    total_duration = (end_time - start_time).total_seconds()
    
    print("\n" + "=" * 80)
    print("Test Results")
    print("=" * 80)
    print(f"Total test duration: {total_duration:.2f} seconds")
    
    if writer_result == 0 and reader_result == 0:
        print("✓ Both writer and reader completed successfully")
        print("\nGenerating Excel report...")
        generate_excel_report()
    else:
        if writer_result != 0:
            print("✗ Writer failed")
        if reader_result != 0:
            print("✗ Reader failed")
    
    print("=" * 80)

def generate_excel_report():
    """Generate Excel report with latency data"""
    try:
        # Load write and read timestamps
        with open('write_timestamps.json', 'r') as f:
            write_data = json.load(f)
        
        with open('read_timestamps.json', 'r') as f:
            read_data = json.load(f)
        
        # Create a mapping of key to write timestamp
        write_map = {item['key']: item for item in write_data}
        
        # Merge data
        merged_data = []
        for read_item in read_data:
            key = read_item['key']
            if key in write_map:
                write_item = write_map[key]
                merged_data.append({
                    'key_number': read_item['key_number'],
                    'key': key,
                    'write_timestamp': write_item['write_timestamp_dt'],
                    'read_timestamp': read_item['read_timestamp_dt'],
                    'latency_ms': read_item['latency_ms'],
                    'latency_seconds': read_item['latency_seconds'],
                    'retry_attempt': read_item['retry_attempt']
                })
        
        # Sort by key number
        merged_data.sort(key=lambda x: x['key_number'])
        
        # Create Excel workbook
        wb = Workbook()
        
        # Sheet 1: Detailed Latency Data
        ws_details = wb.active
        ws_details.title = "Replication Latency"
        
        # Headers
        headers = ["Key Number", "Key", "Write Timestamp (UAE North)", 
                   "Read Timestamp (Sweden Central)", "Latency (ms)", 
                   "Latency (seconds)", "Retry Attempt"]
        ws_details.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws_details[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add data
        for item in merged_data:
            ws_details.append([
                item['key_number'],
                item['key'],
                item['write_timestamp'],
                item['read_timestamp'],
                round(item['latency_ms'], 2),
                round(item['latency_seconds'], 3),
                item['retry_attempt']
            ])
        
        # Adjust column widths
        ws_details.column_dimensions['A'].width = 12
        ws_details.column_dimensions['B'].width = 20
        ws_details.column_dimensions['C'].width = 28
        ws_details.column_dimensions['D'].width = 28
        ws_details.column_dimensions['E'].width = 18
        ws_details.column_dimensions['F'].width = 15
        ws_details.column_dimensions['G'].width = 15
        
        # Sheet 2: Statistics
        ws_stats = wb.create_sheet("Statistics")
        ws_stats.append(["Metric", "Value"])
        
        # Style statistics headers
        for cell in ws_stats[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Calculate statistics (using milliseconds)
        latencies_ms = [item['latency_ms'] for item in merged_data]
        if latencies_ms:
            avg_latency_ms = sum(latencies_ms) / len(latencies_ms)
            min_latency_ms = min(latencies_ms)
            max_latency_ms = max(latencies_ms)
            
            ws_stats.append(["Total Keys", len(merged_data)])
            ws_stats.append(["Average Latency (ms)", round(avg_latency_ms, 2)])
            ws_stats.append(["Minimum Latency (ms)", round(min_latency_ms, 2)])
            ws_stats.append(["Maximum Latency (ms)", round(max_latency_ms, 2)])
            ws_stats.append(["Average Latency (seconds)", round(avg_latency_ms / 1000, 3)])
            ws_stats.append(["Minimum Latency (seconds)", round(min_latency_ms / 1000, 3)])
            ws_stats.append(["Maximum Latency (seconds)", round(max_latency_ms / 1000, 3)])
            ws_stats.append(["Source Region", "UAE North"])
            ws_stats.append(["Destination Region", "Sweden Central"])
            ws_stats.append(["Test Time", datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        
        ws_stats.column_dimensions['A'].width = 30
        ws_stats.column_dimensions['B'].width = 25
        
        # Save workbook
        output_file = f"redis_replication_latency_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(output_file)
        
        print(f"✓ Excel report generated: {output_file}")
        print(f"  - Total keys analyzed: {len(merged_data)}")
        if latencies_ms:
            print(f"  - Average latency: {avg_latency_ms:.2f} ms ({avg_latency_ms/1000:.3f} seconds)")
            print(f"  - Min latency: {min_latency_ms:.2f} ms ({min_latency_ms/1000:.3f} seconds)")
            print(f"  - Max latency: {max_latency_ms:.2f} ms ({max_latency_ms/1000:.3f} seconds)")
    
    except FileNotFoundError as e:
        print(f"ERROR: Could not find timestamp files: {e}")
    except Exception as e:
        print(f"ERROR: Failed to generate Excel report: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    run_replication_test()
