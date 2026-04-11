"""
Run publisher and subscriber in parallel with comprehensive latency tracking
"""
import os
import json
import asyncio
import sys
from datetime import datetime
from azure.eventhub.aio import EventHubProducerClient, EventHubConsumerClient
from azure.eventhub import EventData
from azure.identity.aio import DefaultAzureCredential
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# Configuration
EVENTHUB_NAMESPACE = os.getenv("EVENTHUB_NAMESPACE")
EVENTHUB_NAME = os.getenv("EVENTHUB_NAME")
NUM_EVENTS = int(os.getenv("NUM_EVENTS", 1000))

# Sample event data
SAMPLE_EVENT = {
    "event_type": "test_event",
    "message": "This is a test event for performance measurement",
    "data": {
        "sensor_id": "sensor_001",
        "temperature": 25.5,
        "humidity": 60.0,
        "status": "active"
    }
}

# Global storage for events
published_events = []
received_events = []
events_received_count = 0


async def on_event(partition_context, event):
    """Process received event"""
    global events_received_count
    
    # Record received timestamp
    received_timestamp = datetime.utcnow()
    
    try:
        # Parse event data
        event_data = json.loads(event.body_as_str())
        event_id = event_data.get("event_id", "unknown")
        event_number = event_data.get("event_number", 0)
        
        # Store event info
        received_events.append({
            "event_number": event_number,
            "event_id": event_id,
            "received_timestamp": received_timestamp,
            "partition_id": partition_context.partition_id,
            "sequence_number": event.sequence_number,
            "enqueued_time": event.enqueued_time
        })
        
        events_received_count += 1
        
        # Log progress every 100 events
        if events_received_count % 100 == 0:
            print(f"  📥 Received {events_received_count} events...")
        
        # Update checkpoint
        await partition_context.update_checkpoint(event)
        
    except Exception as e:
        print(f"ERROR processing event: {str(e)}")


async def subscriber_task(stop_event):
    """Subscribe to events from Event Hub"""
    credential = DefaultAzureCredential()
    
    try:
        client = EventHubConsumerClient(
            fully_qualified_namespace=EVENTHUB_NAMESPACE,
            eventhub_name=EVENTHUB_NAME,
            consumer_group="$Default",
            credential=credential
        )
        
        async with client:
            receive_task = asyncio.create_task(
                client.receive(
                    on_event=on_event,
                    starting_position="-1",  # Start from beginning
                )
            )
            
            # Wait for stop signal
            await stop_event.wait()
            
            # Cancel receive task
            receive_task.cancel()
            try:
                await receive_task
            except asyncio.CancelledError:
                pass
                
    except Exception as e:
        print(f"Subscriber error: {str(e)}")


async def publisher_task():
    """Publish events to Event Hub"""
    credential = DefaultAzureCredential()
    
    try:
        producer = EventHubProducerClient(
            fully_qualified_namespace=EVENTHUB_NAMESPACE,
            eventhub_name=EVENTHUB_NAME,
            credential=credential
        )
        
        async with producer:
            print(f"  📤 Publishing {NUM_EVENTS} events...")
            
            for i in range(NUM_EVENTS):
                # Create event with unique ID and timestamp
                event_id = f"event_{i+1}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S_%f')}"
                
                event = SAMPLE_EVENT.copy()
                event["event_id"] = event_id
                event["event_number"] = i + 1
                
                # Create a batch with single event
                batch = await producer.create_batch()
                event_json = json.dumps(event)
                batch.add(EventData(event_json))
                
                # Record publish start time
                publish_start = datetime.utcnow()
                
                # Send batch and wait for acknowledgement
                await producer.send_batch(batch)
                
                # Record acknowledgement received time
                ack_received = datetime.utcnow()
                
                # Calculate latency
                ack_latency_ms = (ack_received - publish_start).total_seconds() * 1000
                
                # Store event info
                published_events.append({
                    "event_number": i + 1,
                    "event_id": event_id,
                    "publish_start": publish_start,
                    "ack_received": ack_received,
                    "ack_latency_ms": ack_latency_ms
                })
                
                # Log progress every 100 events
                if (i + 1) % 100 == 0:
                    avg_latency = sum(e["ack_latency_ms"] for e in published_events) / len(published_events)
                    print(f"  📤 Published {i + 1} events (Avg ACK latency: {avg_latency:.2f}ms)")
            
            print(f"  ✓ All {NUM_EVENTS} events published")
            
    except Exception as e:
        print(f"Publisher error: {str(e)}")


def create_excel_report():
    """Create comprehensive Excel report with all latencies"""
    print("\n📊 Creating Excel report...")
    
    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # Create header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Sheet 1: Complete Event Timeline
    ws_timeline = wb.create_sheet("Event Timeline")
    ws_timeline.append([
        "Event Number",
        "Event ID",
        "Publish Start",
        "ACK Received",
        "Subscriber Read",
        "Publish→ACK (ms)",
        "Publish→Read (ms)",
        "ACK→Read (ms)"
    ])
    
    # Apply header styling
    for cell in ws_timeline[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Match published and received events
    received_dict = {e["event_id"]: e for e in received_events}
    
    timeline_data = []
    for pub_event in published_events:
        event_id = pub_event["event_id"]
        rec_event = received_dict.get(event_id)
        
        if rec_event:
            publish_to_read_ms = (rec_event["received_timestamp"] - pub_event["publish_start"]).total_seconds() * 1000
            ack_to_read_ms = (rec_event["received_timestamp"] - pub_event["ack_received"]).total_seconds() * 1000
            
            timeline_data.append({
                "event_number": pub_event["event_number"],
                "event_id": event_id,
                "publish_start": pub_event["publish_start"].strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
                "ack_received": pub_event["ack_received"].strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
                "subscriber_read": rec_event["received_timestamp"].strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
                "publish_to_ack_ms": round(pub_event["ack_latency_ms"], 2),
                "publish_to_read_ms": round(publish_to_read_ms, 2),
                "ack_to_read_ms": round(ack_to_read_ms, 2)
            })
    
    # Write timeline data
    for data in timeline_data:
        ws_timeline.append([
            data["event_number"],
            data["event_id"],
            data["publish_start"],
            data["ack_received"],
            data["subscriber_read"],
            data["publish_to_ack_ms"],
            data["publish_to_read_ms"],
            data["ack_to_read_ms"]
        ])
    
    # Auto-adjust column widths
    for column in ws_timeline.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws_timeline.column_dimensions[column_letter].width = adjusted_width
    
    # Sheet 2: Statistics
    ws_stats = wb.create_sheet("Statistics")
    ws_stats.append(["Metric", "Value"])
    
    # Apply header styling
    for cell in ws_stats[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Calculate statistics
    if timeline_data:
        publish_to_ack = [d["publish_to_ack_ms"] for d in timeline_data]
        publish_to_read = [d["publish_to_read_ms"] for d in timeline_data]
        ack_to_read = [d["ack_to_read_ms"] for d in timeline_data]
        
        stats = [
            ["Total Events Published", len(published_events)],
            ["Total Events Received", len(received_events)],
            ["Events Matched", len(timeline_data)],
            ["", ""],
            ["Publish → ACK Latency (ms)", ""],
            ["  Average", f"{sum(publish_to_ack)/len(publish_to_ack):.2f}"],
            ["  Minimum", f"{min(publish_to_ack):.2f}"],
            ["  Maximum", f"{max(publish_to_ack):.2f}"],
            ["", ""],
            ["Publish → Subscriber Read Latency (ms)", ""],
            ["  Average", f"{sum(publish_to_read)/len(publish_to_read):.2f}"],
            ["  Minimum", f"{min(publish_to_read):.2f}"],
            ["  Maximum", f"{max(publish_to_read):.2f}"],
            ["", ""],
            ["ACK → Subscriber Read Latency (ms)", ""],
            ["  Average", f"{sum(ack_to_read)/len(ack_to_read):.2f}"],
            ["  Minimum", f"{min(ack_to_read):.2f}"],
            ["  Maximum", f"{max(ack_to_read):.2f}"],
        ]
        
        for row in stats:
            ws_stats.append(row)
    
    # Auto-adjust column widths
    for column in ws_stats.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws_stats.column_dimensions[column_letter].width = adjusted_width
    
    # Save workbook
    filename = f"eventhub_performance_test_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    print(f"✓ Excel report saved: {filename}")
    
    return filename


async def main():
    """Run publisher and subscriber in parallel"""
    
    print("\n" + "=" * 80)
    print("Event Hub Performance Test - Parallel Execution")
    print("=" * 80 + "\n")
    
    if not EVENTHUB_NAMESPACE or not EVENTHUB_NAME:
        print("ERROR: Please set EVENTHUB_NAMESPACE and EVENTHUB_NAME in .env file")
        sys.exit(1)
    
    print(f"Event Hub: {EVENTHUB_NAME}")
    print(f"Namespace: {EVENTHUB_NAMESPACE}")
    print(f"Number of Events: {NUM_EVENTS}")
    print("-" * 80 + "\n")
    
    try:
        # Create stop event for subscriber
        stop_event = asyncio.Event()
        
        # Start subscriber first
        print("🚀 Starting subscriber...")
        subscriber = asyncio.create_task(subscriber_task(stop_event))
        
        # Wait a bit for subscriber to be ready
        await asyncio.sleep(3)
        
        # Start publisher
        print("🚀 Starting publisher...\n")
        await publisher_task()
        
        # Wait a bit for all events to be received
        print("\n⏳ Waiting for all events to be received...")
        await asyncio.sleep(5)
        
        # Stop subscriber
        print("🛑 Stopping subscriber...")
        stop_event.set()
        await subscriber
        
        print("\n" + "=" * 80)
        print(f"✓ Test completed!")
        print(f"  Published: {len(published_events)} events")
        print(f"  Received: {len(received_events)} events")
        print("=" * 80 + "\n")
        
        # Create Excel report
        filename = create_excel_report()
        
        print("\n" + "=" * 80)
        print(f"✓ Report ready: {filename}")
        print("=" * 80 + "\n")
        
    except KeyboardInterrupt:
        print("\n\nTest interrupted by user")
        sys.exit(1)
    except Exception as e:
        print(f"\n\nERROR: Test failed: {str(e)}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
